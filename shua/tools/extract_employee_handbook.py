"""Extract the employee handbook DOCX into browser-friendly data and image assets.

Run with the bundled workspace Python runtime.  The source handbook is read-only;
the generated data file and assets are deterministic build outputs for the app.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path
from typing import Iterator, Union
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


HEADING_LEVELS = {
    "Heading 2": 1,
    "Heading 3": 2,
    "Heading 4": 3,
    "Heading 5": 4,
    "Heading 6": 5,
}


def iter_blocks(document: Document) -> Iterator[Union[Paragraph, Table]]:
    """Yield body paragraphs and tables in the original document order."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def clean_heading(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r"【(?:更新人|撰写人)[^】]*】", "", value)
    return re.sub(r"\s+", " ", value).strip(" -—")


def image_names(paragraph: Paragraph, document: Document) -> list[str]:
    names: list[str] = []
    for element in [*paragraph._p.xpath(".//a:blip"), *paragraph._p.xpath(".//*[local-name()='imagedata']")]:
        relationship_id = element.get(qn("r:embed")) or element.get(qn("r:id"))
        if not relationship_id:
            continue
        part = document.part.related_parts.get(relationship_id)
        if not part:
            continue
        name = Path(str(part.partname)).name
        if name not in names:
            names.append(name)
    return names


def table_block(table: Table, document: Document) -> dict:
    rows = []
    images: list[str] = []
    for row in table.rows:
        row_values = []
        for cell in row.cells:
            value = "\n".join(clean_text(paragraph.text) for paragraph in cell.paragraphs if clean_text(paragraph.text))
            row_values.append(value)
            for paragraph in cell.paragraphs:
                for name in image_names(paragraph, document):
                    if name not in images:
                        images.append(name)
        if any(row_values):
            rows.append(row_values)
    block = {"kind": "table", "rows": rows}
    if images:
        block["images"] = images
    return block


def add_block(target: dict, block: dict) -> None:
    target.setdefault("blocks", []).append(block)


def is_editorial_heading(title: str) -> bool:
    return title.startswith("（") or title.startswith("(")


def split_embedded_heading(title: str) -> tuple[str, str]:
    match = re.match(r"^(备注[：:].*?)(\d+(?:[.,，]\d+)+(?:[.,，]\d+)*.*)$", title)
    return (match.group(1).strip(), match.group(2).strip()) if match else ("", title)


def build_handbook(document: Document) -> dict:
    root = {"title": "新员工上岗手册", "children": [], "blocks": []}
    stack: list[tuple[int, dict]] = [(0, root)]

    for item in iter_blocks(document):
        if isinstance(item, Table):
            add_block(stack[-1][1], table_block(item, document))
            continue

        text = clean_text(item.text)
        images = image_names(item, document)
        level = HEADING_LEVELS.get(item.style.name)
        if level and text:
            title = clean_heading(text)
            note, title = split_embedded_heading(title)
            if note:
                add_block(stack[-1][1], {"kind": "notice", "text": note})
            if is_editorial_heading(title):
                add_block(stack[-1][1], {"kind": "notice", "text": title, "images": images})
                continue
            while stack and stack[-1][0] >= level:
                stack.pop()
            section = {"title": title, "children": [], "blocks": []}
            stack[-1][1]["children"].append(section)
            stack.append((level, section))
            if images:
                add_block(section, {"kind": "images", "items": images})
            continue

        if text or images:
            block = {"kind": "paragraph", "text": text}
            if images:
                block["images"] = images
            add_block(stack[-1][1], block)

    return root


def extract_media(source: Path, asset_dir: Path) -> list[str]:
    if asset_dir.exists():
        shutil.rmtree(asset_dir)
    asset_dir.mkdir(parents=True, exist_ok=True)
    names = []
    with ZipFile(source) as archive:
        for member in archive.namelist():
            if not member.startswith("word/media/") or member.endswith("/"):
                continue
            name = Path(member).name
            (asset_dir / name).write_bytes(archive.read(member))
            names.append(name)
    return sorted(names)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--data-output", type=Path, required=True)
    parser.add_argument("--asset-dir", type=Path, required=True)
    parser.add_argument("--report-output", type=Path, required=True)
    args = parser.parse_args()

    document = Document(args.source)
    handbook = build_handbook(document)
    media = extract_media(args.source, args.asset_dir)
    payload = {
        "version": "2026-02-05",
        "title": "新员工上岗手册",
        "sourceTitle": args.source.stem,
        "mediaBase": "assets/employee-handbook/",
        "content": handbook,
    }
    args.data_output.parent.mkdir(parents=True, exist_ok=True)
    args.data_output.write_text(
        "/* 由 tools/extract_employee_handbook.py 从新员工上岗手册生成，请勿手工编辑。 */\n"
        f"window.EMPLOYEE_HANDBOOK = Object.freeze({json.dumps(payload, ensure_ascii=False, separators=(',', ':'))});\n",
        encoding="utf-8",
    )

    def count_sections(section: dict) -> int:
        return sum(1 + count_sections(child) for child in section.get("children", []))

    def count_blocks(section: dict) -> int:
        return len(section.get("blocks", [])) + sum(count_blocks(child) for child in section.get("children", []))

    def referenced_media(section: dict) -> set[str]:
        names: set[str] = set()
        for block in section.get("blocks", []):
            names.update(block.get("images", []))
            names.update(block.get("items", []))
        for child in section.get("children", []):
            names.update(referenced_media(child))
        return names

    referenced = referenced_media(handbook)
    report = {
        "source": str(args.source),
        "rootCategories": [child["title"] for child in handbook["children"]],
        "sectionCount": count_sections(handbook),
        "blockCount": count_blocks(handbook),
        "paragraphCount": len(document.paragraphs),
        "tableCount": len(document.tables),
        "inlineShapeCount": len(document.inline_shapes),
        "mediaCount": len(media),
        "mediaBytes": sum((args.asset_dir / name).stat().st_size for name in media),
        "referencedMediaCount": len(referenced),
        "unreferencedMedia": sorted(set(media) - referenced),
    }
    args.report_output.parent.mkdir(parents=True, exist_ok=True)
    args.report_output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
