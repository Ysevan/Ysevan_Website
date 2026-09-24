from pathlib import Path

from openpyxl import load_workbook
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\xecat\综合服务经理刷题")
SOURCE = ROOT / "1_（新）综合服务经理岗位准入资格考试题库（20230413）.xlsx"
OUTPUT = ROOT / "综合服务经理岗位准入资格考试题库（20230413）.docx"


def font(run, size=10.5, bold=False, color=None):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("第 ")
    font(r, 9, color=(100, 100, 100))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    r = paragraph.add_run(" 页")
    font(r, 9, color=(100, 100, 100))


wb = load_workbook(SOURCE, data_only=True, read_only=True)
ws = wb["题库"]
rows = list(ws.iter_rows(min_row=2, values_only=True))

questions = []
current = None
for row in rows:
    seq, text, correct, qtype, difficulty = row
    if qtype:
        if current:
            questions.append(current)
        current = {
            "seq": seq,
            "text": text,
            "direct_answer": correct,
            "type": qtype,
            "difficulty": difficulty,
            "options": [],
        }
    elif current and text is not None:
        if current["seq"] is None and seq is not None:
            current["seq"] = seq
        current["options"].append((text, correct == "是"))
if current:
    questions.append(current)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(0)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(70)
title.paragraph_format.space_after = Pt(18)
r = title.add_run("综合服务经理岗位准入资格考试题库")
font(r, 22, True, (31, 78, 121))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("2023年4月13日版")
font(r, 13, False, (90, 90, 90))

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.paragraph_format.space_before = Pt(24)
r = info.add_run(f"共 {len(questions)} 题｜单选题 477 题｜多选题 276 题｜判断题 402 题")
font(r, 10.5, False, (90, 90, 90))

doc.add_page_break()

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = header.add_run("综合服务经理岗位准入资格考试题库")
font(r, 9, False, (120, 120, 120))
add_page_number(sec.footer.paragraphs[0])

last_type = None
type_counts = {"单选题": 0, "多选题": 0, "判断题": 0}
letters = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

for q in questions:
    qtype = str(q["type"])
    if qtype != last_type:
        if last_type is not None:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(qtype)
        font(r, 16, True, (31, 78, 121))
        last_type = qtype

    type_counts[qtype] = type_counts.get(qtype, 0) + 1
    display_no = q["seq"] if q["seq"] is not None else type_counts[qtype]
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    no_split(table.rows[0])

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"{display_no}. ")
    font(r, 10.5, True, (31, 78, 121))
    r = p.add_run(str(q["text"] or ""))
    font(r, 10.5, True)
    r = p.add_run(f"  【{qtype}｜难度：{q['difficulty'] or '未标注'}】")
    font(r, 9, False, (110, 110, 110))

    if qtype == "判断题":
        answer = str(q["direct_answer"] or "")
    else:
        correct_labels = []
        for idx, (option, is_correct) in enumerate(q["options"]):
            op = cell.add_paragraph()
            op.paragraph_format.left_indent = Cm(0.45)
            op.paragraph_format.space_after = Pt(2)
            label = letters[idx] if idx < len(letters) else str(idx + 1)
            rr = op.add_run(f"{label}. {option}")
            font(rr, 10)
            if is_correct:
                correct_labels.append(label)
        answer = "、".join(correct_labels)

    ap = cell.add_paragraph()
    ap.paragraph_format.space_before = Pt(3)
    rr = ap.add_run(f"正确答案：{answer or '未标注'}")
    font(rr, 10, True, (179, 76, 40))
    shade(cell, "F7F9FC")

    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(3)

doc.core_properties.title = "综合服务经理岗位准入资格考试题库（20230413）"
doc.core_properties.subject = "由Excel题库转换为Word格式"
doc.save(OUTPUT)
print(OUTPUT)
print(f"questions={len(questions)} types={type_counts}")
